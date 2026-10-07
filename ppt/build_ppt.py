# -*- coding: utf-8 -*-
"""DeepMesh 5분 발표 덱 — 'deepmesh 발표 템플릿.pptx'를 복제해 만든다.
이미지는 영상에 실제로 쓰인 에셋만 사용(쿠버네티스 아이콘, capture/ 원본, 로고·마스코트).
다이어그램·표는 PowerPoint 네이티브 도형으로 그린다.

  python ppt/build_ppt.py   ->  ppt/DeepMesh_발표.pptx
"""
import copy, random, shutil
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "ppt"
TPL = PPT / "deepmesh 발표 템플릿.pptx"
OUT = PPT / "DeepMesh_발표.pptx"
AS = PPT / "assets"

# 영상에 쓰인 캡처 원본을 덱 폴더로 모아 둔다(재빌드 시 자급)
for f in ["C1_head.png", "C4_glasswing_10000_crop.png", "C5_gcloud_threat_crop.png",
          "k1-off-report.jpeg", "r1_html_compare.png"]:
    if not (AS / f).exists():
        shutil.copy(ROOT / "capture" / f, AS / f)

BLUE, GREEN, INK, MUTE = "005BAA", "00A651", "1A1A1A", "5B6675"
BOX, RED, LRED, LGREEN = "E2EAF3", "DC2626", "FDECEA", "E6F5EC"
BORDER, PALE = "C9D6E6", "F6F8FB"
F = "맑은 고딕"
TBL_STYLE = "{073A0DAA-6AF3-43AB-8588-CEC1D06C72B9}"   # 템플릿 표 스타일
CHROME = {"사각형: 잘린 위쪽 모서리 3", "사각형: 잘린 한쪽 모서리 3", "직사각형 4",
          "Shape 0", "Text 1", "Text 2"}

prs = Presentation(str(TPL))
SRC = list(prs.slides)
# 본문 원본은 9번 슬라이드(초록 라벨·수식 객체 없음). 4번은 주황 라벨+수식이 섞인 예외라 쓰지 않는다.
S_TITLE, S_DIV, S_CONT, S_TABLE, S_QA = SRC[0], SRC[2], SRC[8], SRC[15], SRC[17]
N_ORIG = len(SRC)


# ---------------------------------------------------------------- 템플릿 복제
def dup(src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        old = new._element.cSld.find(qn("p:bg"))
        if old is not None:
            new._element.cSld.remove(old)
        new._element.cSld.insert(0, copy.deepcopy(bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    rid = {}
    for k, rel in src.part.rels.items():
        if rel.is_external or "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        rid[k] = new.part.relate_to(rel.target_part, rel.reltype)
    for node in new.shapes._spTree.iter():
        for a in (qn("r:embed"), qn("r:link"), qn("r:id")):
            if node.get(a) in rid:
                node.set(a, rid[node.get(a)])
    return new


def by_name(sl, name):
    return next(s for s in sl.shapes if s.name == name)


def set_text(shape_or_tf, paras):
    """기존 서식(첫 run)을 살린 채 문단 텍스트만 바꾼다."""
    tf = getattr(shape_or_tf, "text_frame", shape_or_tf)
    ps = list(tf.paragraphs)
    for i, txt in enumerate(paras):
        if i < len(ps):
            p = ps[i]
        else:
            np_ = copy.deepcopy(ps[-1]._p); ps[-1]._p.addnext(np_)
            p = _Paragraph(np_, tf); ps.append(p)
        for br in p._p.findall(qn("a:br")):
            p._p.remove(br)
        runs = p.runs
        r = runs[0] if runs else p.add_run()
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
        r.text = txt
    for p in ps[len(paras):]:
        p._p.getparent().remove(p._p)


def content(label, title):
    """상단 바·라벨·제목·로고만 남기고 본문을 XML 수준에서 비운다(수식 mc:AlternateContent 포함)."""
    sl = dup(S_CONT)
    tree = sl.shapes._spTree
    for el in list(tree):
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        cnv = el.find(".//" + qn("p:cNvPr"))
        keep = cnv is not None and cnv.get("name") in CHROME
        if el.tag == qn("p:pic"):
            off = el.find(".//" + qn("a:off"))
            keep = keep or (off is not None and int(off.get("y")) > Inches(6.9))
        if not keep:
            tree.remove(el)
    used = {v for n in tree.iter() for a in (qn("r:embed"), qn("r:link"), qn("r:id")) if (v := n.get(a))}
    for k, rel in list(sl.part.rels.items()):
        if "image" in rel.reltype and k not in used:
            sl.part.drop_rel(k)
    set_text(by_name(sl, "Text 1"), [label])
    set_text(by_name(sl, "Text 2"), [title])
    return sl


def divider(num, name):
    sl = dup(S_DIV)
    set_text(next(s for s in sl.shapes if s.has_text_frame), [f"Part {num}", name])
    return sl


# ---------------------------------------------------------------- 그리기 도구
def rgb(h):
    return RGBColor.from_string(h)


def _font(run, size, color, bold):
    f = run.font
    f.name = F; f.size = Pt(size); f.bold = bold; f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", F)


def _fill_tf(tf, paras, size, color, bold, align, anchor, space):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, pa in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        runs = pa if isinstance(pa, list) else [(pa, {})]
        for t, o in runs:
            r = p.add_run(); r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold))


def text(sl, x, y, w, h, paras, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_tf(tf, paras, size, color, bold, align, anchor, space)
    return tb


def shape(sl, x, y, w, h, fill=None, line=None, lw=1.0, kind=MSO_SHAPE.ROUNDED_RECTANGLE,
          radius=0.08, dash=False, paras=None, size=12, color=INK, bold=False,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.08, space=0):
    sp = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = MSO_LINE.DASH
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    if paras:
        _fill_tf(tf, paras, size, color, bold, align, anchor, space)
    return sp


def pic(sl, f, x, y, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return sl.shapes.add_picture(str(AS / f), Inches(x), Inches(y), **kw)


def arrow(sl, x1, y1, x2, y2, color=BLUE, w=1.75, dash=False, both=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    if both:
        he = etree.SubElement(ln, qn("a:headEnd")); he.set("type", "triangle"); he.set("w", "med"); he.set("len", "med")
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def notes(sl, t):
    sl.notes_slide.notes_text_frame.text = t


# ================================================================= 슬라이드
# 1. 표지 ------------------------------------------------------------------
s = dup(S_TITLE)
team = next(sh for sh in s.shapes if sh.has_text_frame and "신의철" in sh.text_frame.text)
set_text(team, ["부산대학교 정보컴퓨터공학부", "신의철, 정의진, 이시하   |   지도교수 최윤호"])
notes(s, "[10초] 안녕하세요, 부산대학교 정보컴퓨터공학부 졸업과제 DeepMesh 팀입니다. "
         "KD-CNN 기반 경량 서비스메시로 클라우드 네이티브 침입탐지시스템을 설계·구현했습니다.")

# 2. Part 1 ---------------------------------------------------------------
divider(1, "배경 및 문제")

# 3. BACKGROUND -------------------------------------------------------------
s = content("PART 1 · BACKGROUND", "AI가 해킹을 자동화하는 시대")
def case(x, img, ratio, big, head, sub, src):
    shape(s, x, 1.7, 5.95, 3.95, fill="FFFFFF", line=BORDER)
    iw = 5.55; ih = iw / ratio; band = 1.66
    pic(s, img, x + 0.2, 1.9 + (band - ih) / 2, w=iw)
    text(s, x + 0.25, 3.75, 2.75, 0.75, [big], size=38, color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 3.0, 3.8, 2.8, 0.7, head, size=14.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, 4.62, 5.45, 0.55, [sub], size=12.5, color=MUTE)
    text(s, x + 0.25, 5.22, 5.45, 0.3, [src], size=10.5, color=MUTE)
case(0.62, "C1_head.png", 1843 / 551, "80~90%", ["해킹 전 과정 중", "AI가 스스로 수행한 비율"],
     "정찰·침투부터 옆 시스템으로 번지는 측면이동까지 AI가 직접 수행",
     "출처: Anthropic, 2025.11.13")
case(6.77, "C4_glasswing_10000_crop.png", 960 / 237, "10,000+", ["AI로 한 달여 만에 찾은", "심각한 취약점 수"],
     "약 50개 기관 참여 — 공격자도 같은 속도로 취약점을 찾을 수 있다",
     "출처: Anthropic Project Glasswing, 2026.05.22")
shape(s, 0.62, 5.85, 12.1, 0.85, fill=BOX, radius=0.12, paras=[[
    ("결국 언젠가는 뚫린다.   ", {"bold": True}),
    ("→ 침입 이후 내부로 ‘번지는 것’을 막아야 한다", {"bold": True, "color": BLUE})]], size=18)
notes(s, "[35초] 2025년 11월 앤트로픽 보고서에 따르면 한 해킹 공격에서 전체 과정의 80~90%를 AI가 스스로 해냈습니다. "
         "침투는 물론 안으로 들어간 뒤 옆 시스템으로 번지는 것까지요. 2026년에는 약 50개 기관이 AI로 한 달여 만에 "
         "심각한 취약점을 만 개 넘게 찾았습니다. 공격자도 같은 속도로 찾을 수 있으니 결국 언젠가는 뚫린다고 봐야 합니다.")

# 4. PROBLEM ----------------------------------------------------------------
s = content("PART 1 · PROBLEM", "뚫린 뒤에는 내부에서 번진다 — 측면이동")
text(s, 0.62, 1.68, 5.3, 0.35, ["공격자가 처음 들어온 경로"], size=15, bold=True)
pic(s, "C5_gcloud_threat_crop.png", 0.62, 2.08, w=5.2)
for i, (v, l) in enumerate([("44.5%", "SW 취약점"), ("27.2%", "인증정보 부실"), ("21.0%", "설정 실수")]):
    shape(s, 0.62 + i * 1.76, 5.62, 1.68, 0.95, fill=BOX, paras=[
        [(v, {"size": 20, "bold": True, "color": BLUE})], [(l, {"size": 12})]])
text(s, 0.62, 6.66, 5.4, 0.28, ["출처: Google Cloud Threat Horizons Report H1 2026 (2025년 하반기)"], size=10, color=MUTE)
# 클러스터 다이어그램
shape(s, 6.25, 1.95, 6.47, 3.45, fill=PALE, line="8A96A8", dash=True, radius=0.04)
text(s, 6.45, 2.05, 3, 0.3, ["쿠버네티스 클러스터"], size=12, bold=True, color=MUTE)
shape(s, 8.15, 1.76, 3.2, 0.4, fill=LGREEN, line=GREEN, radius=0.5,
      paras=["경계 보안(방화벽): 외부 출입만 검사"], size=11.5, color=GREEN, bold=True)
shape(s, 6.45, 2.95, 1.2, 1.45, fill=LRED, line=RED, lw=1.5)
pic(s, "pod.png", 6.66, 3.02, w=0.78)
text(s, 6.45, 3.83, 1.2, 0.3, ["auth pod"], size=12, bold=True, color=RED, align=PP_ALIGN.CENTER)
shape(s, 6.62, 2.58, 0.86, 0.3, fill=RED, radius=0.5, paras=["탈취"], size=11, color="FFFFFF", bold=True)
pic(s, "api.png", 10.95, 2.3, w=0.95)
text(s, 10.6, 3.27, 1.65, 0.3, ["K8s API 서버"], size=12, bold=True, align=PP_ALIGN.CENTER)
pods = [("frontend", 8.45), ("post", 9.75), ("comment", 11.05)]
for n, x in pods:
    pic(s, "pod.png", x, 4.08, w=0.55)
    text(s, x - 0.33, 4.66, 1.2, 0.3, [n], size=12, align=PP_ALIGN.CENTER)
arrow(s, 7.65, 3.45, 10.95, 2.78, color=RED, w=2)
pic(s, "sa.png", 7.75, 2.5, w=0.36)
text(s, 8.15, 2.55, 2.0, 0.28, ["① SA 토큰으로 정찰"], size=11.5, bold=True, color=RED)
for n, x in pods:
    arrow(s, 7.65, 3.8, x + 0.27, 4.08, color=RED, w=1.75, dash=True)
text(s, 8.45, 3.45, 2.9, 0.28, ["② 다른 pod로 번짐 (측면이동)"], size=11.5, bold=True, color=RED)
shape(s, 6.25, 5.6, 6.47, 1.3, fill=BOX, align=PP_ALIGN.LEFT, margin=0.2, space=4, paras=[
    [("• pod 하나만 탈취돼도 기본 탑재된 서비스계정(SA) 토큰으로 API 서버를 정찰하고 다른 pod로 번진다.", {})],
    [("• 경계 보안은 외부 출입만 본다 → 클러스터 내부 동서(East-West) 트래픽은 ", {}),
     ("무방비", {"bold": True, "color": RED})]], size=12.5)
notes(s, "[40초] 정문은 생각보다 쉽게 뚫립니다. 구글 클라우드 보고서에 따르면 초기 침투 경로의 44.5%가 소프트웨어 취약점, "
         "27.2%가 약한 인증정보, 21%가 설정 실수였습니다. 쿠버네티스에서 pod 하나만 탈취돼도 기본 서비스계정 토큰으로 "
         "API 서버를 정찰하고 옆 pod로 번집니다. 방화벽은 외부 출입만 보기 때문에 내부 트래픽은 무방비입니다.")

# 5. Part 2 ---------------------------------------------------------------
divider(2, "시스템 설계")

# 6. ARCHITECTURE -----------------------------------------------------------
s = content("PART 2 · ARCHITECTURE", "시스템 아키텍처 — 모든 pod에 AI 검사기를 붙인다")
shape(s, 0.62, 1.65, 12.1, 4.2, fill=PALE, line="8A96A8", dash=True, radius=0.03)
pic(s, "k8s.png", 0.78, 1.72, w=0.3)
text(s, 1.15, 1.74, 9, 0.3, ["쿠버네티스 클러스터 · VM 4대 (마스터 1 + 워커 3) · Kubernetes v1.33"], size=12, bold=True, color=MUTE)
shape(s, 0.9, 2.15, 3.25, 3.45, fill="FFFFFF", line=BLUE, lw=1.25, radius=0.05)
pic(s, "cp.png", 1.05, 2.27, w=0.5)
text(s, 1.62, 2.33, 2.5, 0.4, ["마스터 노드 · Control Plane"], size=13, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
for y, h, b in [(2.95, "Pod Info Provider", "형제 replica 주소록을 사이드카에 미리 push"),
                (4.18, "Request Verifier", "이상 요청 시그니처가 형제 replica에서 관측된 적 있는지 조회")]:
    shape(s, 1.05, y, 2.95, 1.1, fill=BOX, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14, space=2,
          paras=[[(h, {"bold": True, "size": 13})], [(b, {"size": 11, "color": MUTE})]])
text(s, 1.05, 5.33, 3, 0.25, ["마스터 노드의 Python 프로그램 (Pod 아님)"], size=10, color=MUTE)
text(s, 4.55, 2.12, 7.9, 0.3, ["워커 노드 — 게시판 서비스 4개 (각 replica 2개)"], size=12, bold=True, color=MUTE)
for i, n in enumerate(["frontend", "auth", "post", "comment"]):
    x = 4.55 + i * 2.0
    shape(s, x, 2.48, 1.9, 3.12, fill="FFFFFF", line=BORDER, radius=0.05)
    text(s, x, 2.56, 1.9, 0.32, [n], size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    for j in range(2):
        y = 2.95 + j * 1.3
        shape(s, x + 0.12, y, 1.66, 1.18, fill=PALE, line="D6E0EC", radius=0.06)
        pic(s, "pod.png", x + 0.2, y + 0.36, w=0.45)
        shape(s, x + 0.72, y + 0.14, 0.94, 0.4, fill="FFFFFF", line=BORDER, radius=0.15,
              paras=["App"], size=10.5)
        shape(s, x + 0.72, y + 0.64, 0.94, 0.4, fill=GREEN, radius=0.15,
              paras=["사이드카"], size=10.5, color="FFFFFF", bold=True)
arrow(s, 4.17, 3.85, 4.53, 3.85, color=BLUE, w=1.5, both=True)
for i, t in enumerate(["① iptables(InitContainer)로 pod의 outbound 트래픽을 사이드카로 가로챔",
                       "② 사이드카 안에서 가로채기 → 이미지화 → 판정",
                       "③ 이상 시 요청은 DROP, 응답은 형제 replica 응답으로 RELAY"]):
    shape(s, 0.62 + i * 4.075, 6.0, 3.95, 0.85, fill=BOX, paras=[t], size=12.5,
          align=PP_ALIGN.LEFT, margin=0.15)
notes(s, "[45초] 실험 환경은 VM 4대로 만든 쿠버네티스 클러스터입니다. 게시판을 frontend, auth, post, comment 네 서비스로 나누고 "
         "각각 replica 2개를 띄웠습니다. 모든 pod에 사이드카를 붙여 iptables로 나가는 트래픽을 가로채 검사합니다. "
         "마스터 노드의 Control Plane은 형제 replica 주소록을 미리 나눠 주고, 이상 요청이 형제에게서도 관측됐는지 확인합니다.")

# 7. DETECTION --------------------------------------------------------------
s = content("PART 2 · DETECTION", "탐지 파이프라인 — 정상만 배우고, 작게 줄인다")
steps = [("가로채기", "사이드카가 outbound 패킷을 출발지→목적지 세션 단위로 묶음"),
         ("이미지화", "패킷 5개 × 의미 특징 20개 → 20×5 흑백 이미지"),
         ("KD-CNN", "경량 학생 모델이 이미지를 128차원 임베딩으로 변환"),
         ("OCSVM 판정", "정상 임베딩의 경계 밖이면 이상으로 판정"),
         ("대응", None)]
for i, (h, b) in enumerate(steps):
    x = 0.62 + i * 2.475
    last = i == 4
    shape(s, x, 1.7, 2.2, 2.05, fill="FFFFFF" if last else BOX, line=GREEN if last else None, lw=1.5)
    shape(s, x + 0.15, 1.85, 0.42, 0.42, kind=MSO_SHAPE.OVAL, fill=GREEN if last else BLUE,
          paras=[str(i + 1)], size=14, color="FFFFFF", bold=True, margin=0)
    text(s, x + 0.65, 1.85, 1.5, 0.42, [h], size=14.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if last:
        text(s, x + 0.15, 2.45, 1.95, 1.2, [
            [("정상 → ", {}), ("FORWARD", {"bold": True, "color": GREEN})],
            [("요청 이상 → ", {}), ("DROP", {"bold": True, "color": RED})],
            [("응답 이상 → ", {}), ("RELAY", {"bold": True, "color": BLUE})]], size=12, space=6)
    elif i == 1:
        text(s, x + 0.15, 2.42, 1.25, 1.25, [b], size=11.5)
        rnd = random.Random(7)
        for r in range(20):
            for c in range(5):
                g = rnd.choice([0x22, 0x55, 0x88, 0xBB, 0xEE, 0xFF])
                shape(s, x + 1.5 + c * 0.11, 2.42 + r * 0.058, 0.11, 0.058, fill=f"{g:02X}{g:02X}{g:02X}",
                      kind=MSO_SHAPE.RECTANGLE)
        shape(s, x + 1.5, 2.42, 0.55, 1.16, line="8A96A8", lw=0.75, kind=MSO_SHAPE.RECTANGLE)
    else:
        text(s, x + 0.15, 2.42, 1.95, 1.25, [b], size=11.5)
    if not last:
        arrow(s, x + 2.22, 2.72, x + 2.455, 2.72, color=BLUE, w=2)
# 비지도 학습
shape(s, 0.62, 4.0, 5.95, 2.85, fill="FFFFFF", line=BORDER, radius=0.05)
text(s, 0.85, 4.12, 5.5, 0.4, ["정상 트래픽만으로 학습 (비지도)"], size=15, bold=True, color=BLUE)
shape(s, 0.95, 4.7, 2.2, 1.6, kind=MSO_SHAPE.OVAL, fill=LGREEN, line=GREEN, lw=1.5, dash=True)
rnd = random.Random(3)
for _ in range(12):
    shape(s, 1.35 + rnd.random() * 1.3, 5.0 + rnd.random() * 0.95, 0.11, 0.11, kind=MSO_SHAPE.OVAL, fill=GREEN)
text(s, 1.2, 6.33, 1.7, 0.28, ["정상 경계"], size=11, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
shape(s, 3.22, 4.62, 0.17, 0.17, kind=MSO_SHAPE.OVAL, fill=RED)
text(s, 3.02, 4.36, 0.6, 0.26, ["이상"], size=11, bold=True, color=RED, align=PP_ALIGN.CENTER)
text(s, 3.65, 4.72, 2.75, 1.6, [
    "공격 데이터 없이 정상의 경계를 학습 → 처음 보는 공격도 ‘정상과 다르면’ 탐지",
    [("※ 공격 이미지는 임계값 재보정에만 사용", {"size": 10.5, "color": MUTE})]], size=12.5, space=8)
# 지식증류
shape(s, 6.77, 4.0, 5.95, 2.85, fill="FFFFFF", line=BORDER, radius=0.05)
text(s, 7.0, 4.12, 5.5, 0.4, ["지식증류로 경량화 (Teacher → Student)"], size=15, bold=True, color=BLUE)
shape(s, 7.0, 4.62, 1.75, 1.75, kind=MSO_SHAPE.OVAL, fill="DCE6F2", line=BLUE, lw=1.5,
      paras=[[("Teacher", {"size": 12})], [("314.69K", {"size": 15, "bold": True})]], color=BLUE)
arrow(s, 8.9, 5.5, 9.55, 5.5, color=GREEN, w=2)
d = 1.75 * (12.64 / 314.69) ** 0.5
shape(s, 9.95 - d / 2, 5.5 - d / 2, d, d, kind=MSO_SHAPE.OVAL, fill=GREEN)
text(s, 9.35, 5.82, 1.2, 0.6, [[("Student", {"size": 11.5})], [("1.2K~12.6K", {"size": 13, "bold": True})]],
     color=GREEN, align=PP_ALIGN.CENTER)
text(s, 10.75, 4.75, 1.85, 1.8, [[("파라미터", {})], [("약 25~256배↓", {"bold": True, "color": GREEN})],
     [("GPU 없이 CPU로", {})], [("판정 약 1ms", {"bold": True, "color": BLUE})]], size=12.5, space=3)
notes(s, "[45초] 사이드카는 나가는 패킷을 세션으로 묶고, 패킷 5개의 의미 특징 20개를 20×5 이미지로 바꿉니다. "
         "경량 CNN이 이를 임베딩으로 만들고 OCSVM이 정상 경계 밖인지 판정합니다. 정상 트래픽만으로 학습해 처음 보는 공격도 잡고, "
         "큰 Teacher 모델을 지식증류로 작은 Student로 옮겨 31만 개 파라미터를 천~만 개 수준으로 줄였습니다.")

# 8. Part 3 ---------------------------------------------------------------
divider(3, "시연 및 결과")

# 9. DEMO -------------------------------------------------------------------
s = content("PART 3 · DEMO", "공격 시연 — 사이드카 OFF vs ON")
def scenario(x, tag, title, img, ratio, cap, off, on_head, on, boxes=None, src_w=None):
    shape(s, x, 1.68, 1.75, 0.34, fill=BLUE, radius=0.5, paras=[tag], size=11.5, color="FFFFFF", bold=True)
    text(s, x, 2.08, 5.95, 0.45, [title], size=15.5, bold=True)
    iw = 3.3
    pic(s, img, x, 2.7, w=iw)
    for bx, by, bw, bh in (boxes or []):
        k = iw / src_w
        shape(s, x + bx * k, 2.7 + by * k, bw * k, bh * k, line=RED, lw=1.75, kind=MSO_SHAPE.RECTANGLE)
    text(s, x, 2.7 + iw / ratio + 0.06, iw, 0.28, [cap], size=10, color=MUTE)
    shape(s, x + 3.45, 2.7, 2.5, 1.3, fill=LRED, line=RED, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          margin=0.14, space=3, paras=[[("사이드카 OFF", {"bold": True, "color": RED, "size": 12.5})], off], size=11.5)
    shape(s, x + 3.45, 4.15, 2.5, 1.34, fill=LGREEN, line=GREEN, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          margin=0.14, space=3, paras=[[(on_head, {"bold": True, "color": GREEN, "size": 12.5})], on], size=11.5)
scenario(0.62, "시나리오 1 · k1", "탈취한 auth pod → 쿠버네티스 API 정찰",
         "k1-off-report.jpeg", 1000 / 844, "사이드카 OFF 시 정찰 결과 (실제 캡처)",
         "API 서버가 200으로 응답 → 클러스터 버전·정보 노출",
         "사이드카 ON → DROP", "처음 보는 방향(API 서버)으로 나가는 요청을 이상 판정, 연결 전에 차단",
         boxes=[(24, 280, 380, 30), (8, 398, 515, 292)], src_w=1000)
scenario(6.77, "시나리오 2 · r1", "frontend 응답에 XSS 스크립트 주입",
         "r1_html_compare.png", 1500 / 1208, "위: 변조 응답(OFF) · 아래: 교체된 정상 응답(ON)",
         "변조 응답이 그대로 전달 → 브라우저에서 스크립트 실행, attacker.test로 요청 발생",
         "사이드카 ON → RELAY", "형제 replica에 같은 요청을 보내 응답을 비교 → 정상 응답으로 교체")
shape(s, 0.62, 5.88, 12.1, 0.78, fill=BOX, radius=0.12, paras=[
    "실제 쿠버네티스 클러스터(VM 4대)와 자체 게시판에서 재현 · 관리자 대시보드로 탐지 결과를 실시간 확인"], size=13.5)
notes(s, "[45초] 첫 번째는 auth pod를 탈취해 서비스계정 토큰으로 API 서버를 정찰하는 공격입니다. 사이드카가 꺼져 있으면 "
         "API 서버가 그대로 응답해 클러스터 정보가 노출되지만, 켜면 처음 보는 방향의 요청으로 잡아 차단합니다. "
         "두 번째는 frontend 응답에 XSS 스크립트를 끼워 넣는 공격입니다. 사이드카는 응답이 평소와 다르다는 걸 탐지하고, "
         "형제 replica의 응답과 비교해 정상 응답으로 바꿔서 내보냅니다.")

# 10. RESULT ----------------------------------------------------------------
s = content("PART 3 · RESULT", "성능 평가 — 정확하고, 가볍고, 빠르다")
for i, (v, l, c) in enumerate([("86.7~100%", "측면이동 공격 재현율 (Recall)", GREEN),
                               ("≤ 1.3%", "정상을 공격으로 본 오탐률 (FPR)", BLUE),
                               ("25~256배↓", "파라미터 314.69K → 1.2~12.6K", GREEN),
                               ("0.36~1.13 ms", "CPU 이미지 1장 판정 시간", BLUE)]):
    shape(s, 0.62 + i * 3.056, 1.7, 2.93, 1.35, fill=BOX, space=4,
          paras=[[(v, {"size": 25, "bold": True, "color": c})], [(l, {"size": 12})]])
rows = [["서비스", "배포 모델", "파라미터", "재현율", "오탐률 (FPR)", "CPU 판정"],
        ["auth", "1x16", "12.64K", "86.67%", "0.32%", "0.38 ms"],
        ["post", "1x8", "1.23K", "96.29%", "0.00%", "0.36 ms"],
        ["comment", "2x8", "5.69K", "95.54%", "1.30%", "1.13 ms"],
        ["frontend*", "2x8", "5.69K", "100.00%", "0.93%", "0.80 ms"],
        ["Teacher (비교)", "—", "314.69K", "—", "—", "—"]]
cw = [2.0, 1.6, 1.9, 2.0, 2.2, 2.4]
tshp = s.shapes.add_table(len(rows), len(cw), Inches(0.62), Inches(3.25), Inches(sum(cw)), Inches(0.5 * len(rows)))
tbl = tshp.table
tblPr = tbl._tbl.tblPr
for a in ("firstRow", "bandRow"):
    tblPr.attrib.pop(a, None)
tblPr.find(qn("a:tableStyleId")).text = TBL_STYLE
for j, w in enumerate(cw):
    tbl.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.5)
    for j, v in enumerate(row):
        cell = tbl.cell(i, j)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        col = INK; bold = i == 0
        if i > 0 and j == 3 and v != "—": col, bold = GREEN, True
        if i > 0 and j == 2: col, bold = BLUE, True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run(); r.text = v
        _font(r, 13, col, bold)
text(s, 0.62, 6.34, 12.1, 0.55, [
    "평가: CPU · batch 1 · 스레드 1 오프라인 평가. 판정 시간은 모델 추론(텐서 변환 + CNN + OCSVM)만 측정.",
    "* frontend는 검증된 운영 임계값이 아닌 학습 시점의 잠정값.  Teacher는 비교 기준이며 배포하지 않음."], size=10, color=MUTE)
notes(s, "[35초] 정상 트래픽만으로 학습한 모델이 측면이동 공격을 서비스별 재현율 86.7~100%로 탐지했고, "
         "정상을 공격으로 잘못 본 비율은 1.3% 이하였습니다. 파라미터는 31만 개에서 천~만 개 수준으로 줄었고, "
         "GPU 없이 CPU만으로 이미지 한 장을 약 1ms 안팎에 판정합니다.")

# 11. CONCLUSION (템플릿 표 슬라이드 복제) ---------------------------------------
s = dup(S_TABLE)
set_text(by_name(s, "Text 1"), ["PART 3 · CONCLUSION"])
set_text(by_name(s, "Text 2"), ["결론 및 향후 과제"])
t = next(sh for sh in s.shapes if sh.has_table).table
data = [["#", "구분", "내용", "의의"],
        ["1", "서비스메시 기반 IDS", "모든 pod의 사이드카가 outbound 트래픽을 전수 검사", "내부 측면이동 실시간 차단"],
        ["2", "비지도 학습", "정상 트래픽만으로 OCSVM 경계 학습", "처음 보는 공격 대응"],
        ["3", "지식증류 경량화", "314.69K → 1.2~12.6K, CPU만으로 판정 약 1ms", "사이드카 상주 가능"],
        ["4", "한계", "형제 replica가 정상이라는 전제 · RELAY는 replica 1개와 비교", "3개 이상 다수결로 확장"],
        ["5", "향후 과제", "frontend 운영 임계값 검증, 실서비스 규모로 확장", "실운영 검증"]]
for i, row in enumerate(data):
    for j, v in enumerate(row):
        cell = t.cell(i, j)
        set_text(cell.text_frame, [v])
        if i > 0 and j == 3:
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.color.rgb = rgb(BLUE if i <= 3 else GREEN); run.font.bold = True
notes(s, "[25초] 정리하면, 서비스메시의 사이드카에 경량 AI 검사기를 붙여 내부 측면이동을 실시간으로 막았습니다. "
         "한계로는 형제 replica가 정상이라는 전제가 있어, 3개 이상의 replica 다수결로 확장할 계획입니다. 감사합니다.")

# 12. Q&A ---------------------------------------------------------------------
dup(S_QA)

# ---------------------------------------------------------------- 원본 18장 제거
lst = prs.slides._sldIdLst
for sid in list(lst)[:N_ORIG]:
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)

prs.save(str(OUT))
print("saved", OUT, "slides", len(prs.slides))
